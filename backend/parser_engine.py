import os
import re
import unicodedata
import json
import mimetypes
from typing import Any, Dict, List, Optional, Tuple

import google.generativeai as genai
import openpyxl
import pandas as pd
import xlrd
from dotenv import load_dotenv

from backend.database import get_db_connection
import backend.supabase_store as supabase_store


def _supabase_strict_enabled() -> bool:
    return getattr(supabase_store, "strict_enabled", lambda: True)()


load_dotenv()
api_key = os.getenv("GEMINI_API_KEY") or os.getenv("GOOGLE_API_KEY")
if api_key:
    genai.configure(api_key=api_key)


IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}


def _gemini_model(model_name: str):
    key = os.getenv("GEMINI_API_KEY") or os.getenv("GOOGLE_API_KEY")
    if not key:
        raise ValueError("Đọc ảnh cần cấu hình GEMINI_API_KEY trong file .env.")
    genai.configure(api_key=key)
    return genai.GenerativeModel(model_name)


def _gemini_model_candidates() -> List[str]:
    configured = os.getenv("GEMINI_MODEL")
    candidates = [configured, "gemini-2.0-flash", "gemini-2.5-flash", "gemini-1.5-flash"]
    seen = set()
    result = []
    for name in candidates:
        if not name or name in seen:
            continue
        seen.add(name)
        result.append(name)
    return result


def _generate_gemini_content(contents: Any, generation_config: Optional[Dict[str, Any]] = None):
    errors = []
    for model_name in _gemini_model_candidates():
        try:
            return _gemini_model(model_name).generate_content(contents, generation_config=generation_config)
        except Exception as exc:
            message = str(exc)
            message_lower = message.lower()
            if (
                "resourceexhausted" in message_lower
                or "quota exceeded" in message_lower
                or "429" in message
            ):
                raise RuntimeError(
                    "Gemini API đã hết quota hoặc bị giới hạn tốc độ. "
                    "Hãy chờ quota reset, nâng gói/bật billing, hoặc đổi GEMINI_API_KEY/GEMINI_MODEL trong .env."
                ) from exc
            errors.append(f"{model_name}: {type(exc).__name__}: {message}")
            if "not found" not in message_lower and "not supported" not in message_lower and "404" not in message:
                raise
    raise RuntimeError("Không gọi được Gemini OCR/phân loại với các model đã thử: " + " | ".join(errors))


def normalize_text(value: Any) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.replace("đ", "d").replace("Đ", "D")
    text = text.replace("đ", "d").replace("Đ", "D")
    return re.sub(r"\s+", " ", text.lower()).strip()


def parse_number(value: Any) -> Optional[float]:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", ".")
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    return float(match.group(0)) if match else None


def resolve_merged_cells(sheet) -> Dict[tuple, Any]:
    merged_lookup = {}
    for merged_range in sheet.merged_cells.ranges:
        top_left_val = sheet.cell(row=merged_range.min_row, column=merged_range.min_col).value
        for row in range(merged_range.min_row, merged_range.max_row + 1):
            for col in range(merged_range.min_col, merged_range.max_col + 1):
                merged_lookup[(row, col)] = top_left_val
    return merged_lookup


def detect_headers(sheet, merged_lookup: Dict[tuple, Any], max_scan_rows: int = 35) -> tuple[int, Dict[str, int]]:
    field_keywords = {
        "item_no": ["item", "no", "stt", "tt", "so tt"],
        "code": ["code", "ma", "ma hieu", "ma hieu don gia"],
        "mark": ["mark", "ky hieu", "tag", "symbol", "ky hieu ban ve", "khbv"],
        "description": [
            "description",
            "mo ta",
            "mo ta chi tiet",
            "mo ta hang hoa",
            "dien giai",
            "noi dung",
            "noi dung cong viec",
            "hang muc cong viec",
            "hang muc",
            "ten cong tac",
            "cong tac",
            "danh muc cong tac",
            "ten vat tu",
            "ten vat tu, thiet bi",
            "quy cach",
            "ten san pham",
        ],
        "material": ["material", "vat lieu", "chat lieu", "thong so ky thuat"],
        "pressure_class": ["pressure", "class", "ap suat", "cap ap suat", "ghcl", "ei"],
        "quantity": ["qty", "quantity", "so luong", "sl", "khoi luong", "kl"],
        "unit": ["unit", "don vi", "don vi tinh", "dvt"],
        "remark": ["remark", "note", "ghi chu", "chu thich"],
    }

    best_row = 1
    best_score = 0
    best_mapping: Dict[str, int] = {}

    for row in range(1, min(max_scan_rows, sheet.max_row) + 1):
        values = []
        for col in range(1, sheet.max_column + 1):
            value = merged_lookup.get((row, col), sheet.cell(row=row, column=col).value)
            values.append(normalize_text(value))

        mapping: Dict[str, int] = {}
        score = 0
        for field, keywords in field_keywords.items():
            for idx, cell_value in enumerate(values):
                if any(keyword in cell_value for keyword in keywords):
                    mapping[field] = idx + 1
                    score += 1
                    break

        description_col = mapping.get("description")
        if description_col:
            description_header = values[description_col - 1]
            product_header_col = next(
                (
                    idx + 1
                    for idx, cell_value in enumerate(values)
                    if any(
                        token in cell_value
                        for token in ["san pham", "ten san pham", "ten vat tu", "ten vat tu, thiet bi"]
                    )
                ),
                None,
            )
            if product_header_col and "quy cach" in description_header:
                mapping["description"] = product_header_col
                if "material" not in mapping:
                    mapping["material"] = description_col

        if "description" in mapping and "quantity" in mapping:
            score += 2
        if "unit" in mapping:
            score += 1

        if score > best_score:
            best_row = row
            best_score = score
            best_mapping = mapping

    if best_score < 3:
        return 1, {
            "item_no": 1,
            "code": 2,
            "description": 3,
            "unit": 4,
            "quantity": 5,
        }

    return best_row, best_mapping


def select_takeoff_sheet(workbook) -> Any:
    for sheet in workbook.worksheets:
        if normalize_text(sheet.title) == "kl":
            return sheet

    for sheet in workbook.worksheets:
        title = normalize_text(sheet.title)
        if title.startswith("dnvt") or "de nghi vat tu" in title:
            return sheet

    for sheet in workbook.worksheets:
        title = normalize_text(sheet.title)
        if title == "bao gia":
            continue
        merged_lookup = resolve_merged_cells(sheet)
        sample_text = " ".join(
            normalize_text(merged_lookup.get((row, col), sheet.cell(row=row, column=col).value))
            for row in range(1, min(sheet.max_row, 20) + 1)
            for col in range(1, min(sheet.max_column, 10) + 1)
        )
        if "giay de nghi vat tu" in sample_text:
            return sheet

    best_sheet = workbook.active
    best_score = -1
    for sheet in workbook.worksheets:
        merged_lookup = resolve_merged_cells(sheet)
        _, mapping = detect_headers(sheet, merged_lookup)
        score = len(mapping)
        if "description" in mapping and "quantity" in mapping:
            score += 3
        if score > best_score:
            best_sheet = sheet
            best_score = score
    return best_sheet


def _looks_like_no_header_quantity_matrix(sheet) -> bool:
    hits = 0
    for row in range(1, min(sheet.max_row, 120) + 1):
        description = str(sheet.cell(row=row, column=2).value or "").strip()
        unit = normalize_text(sheet.cell(row=row, column=3).value)
        qty = _quantity_matrix_row_total(sheet, row)
        if description and unit in {"cai", "bo", "m", "m2", "tan", "lo"} and qty > 0:
            hits += 1
    return hits >= 5


def _quantity_matrix_row_total(sheet, row: int) -> float:
    total_cell = sheet.cell(row=row, column=5).value
    total = parse_number(total_cell)
    if total is not None and total > 0:
        return total

    total = 0.0
    for col in range(6, min(sheet.max_column, 12) + 1):
        value = parse_number(sheet.cell(row=row, column=col).value)
        if value:
            total += value
    return total


def _parse_no_header_quantity_matrix(sheet) -> List[Dict[str, Any]]:
    parsed_items = []
    item_no = 1
    for row in range(1, sheet.max_row + 1):
        description = str(sheet.cell(row=row, column=2).value or "").strip()
        unit = str(sheet.cell(row=row, column=3).value or "").strip()
        if not description or normalize_text(unit) not in {"cai", "bo", "m", "m2", "tan", "lo"}:
            continue

        qty = _quantity_matrix_row_total(sheet, row)
        if qty <= 0:
            continue

        material_or_note = str(sheet.cell(row=row, column=4).value or "").strip()
        item = _build_takeoff_item(
            {
                "item_no": item_no,
                "description": description,
                "unit": unit,
                "quantity": qty,
                "material": material_or_note,
                "remark": material_or_note,
            },
            source_row=row,
        )
        if item:
            item["warnings"].append(
                "Dòng đọc từ bảng không có tiêu đề chuẩn; số lượng được cộng từ các cột F:L. Cần kiểm tra lại trước khi phát hành."
            )
            parsed_items.append(item)
            item_no += 1

    parsed_items = assign_marks(parsed_items)
    parsed_items = apply_thickness_rules(parsed_items)
    return parsed_items


def _looks_like_ycbg_request_sheet(sheet) -> bool:
    """Detect simple customer RFQ sheets: A=STT, B=description, D=unit, E=quantity."""
    title_text = normalize_text(sheet.title)
    top_text = " ".join(
        normalize_text(sheet.cell(row=row, column=col).value)
        for row in range(1, min(sheet.max_row, 12) + 1)
        for col in range(1, min(sheet.max_column, 8) + 1)
    )
    if "yeu cau bao gia" not in f"{title_text} {top_text}" and "ycbg" not in f"{title_text} {top_text}":
        return False

    hits = 0
    for row in range(1, min(sheet.max_row, 120) + 1):
        item_no = parse_number(sheet.cell(row=row, column=1).value)
        description = str(sheet.cell(row=row, column=2).value or "").strip()
        unit = normalize_text(sheet.cell(row=row, column=4).value)
        qty = parse_number(sheet.cell(row=row, column=5).value)
        if item_no is not None and description and unit in {"m", "m2", "cai", "bo", "set", "tan", "lo"} and qty and qty > 0:
            hits += 1
    return hits >= 3


def _parse_ycbg_request_sheet(sheet) -> List[Dict[str, Any]]:
    parsed_items: List[Dict[str, Any]] = []
    current_section = ""
    for row in range(1, sheet.max_row + 1):
        first_cell = sheet.cell(row=row, column=1).value
        description = str(sheet.cell(row=row, column=2).value or "").strip()
        unit = str(sheet.cell(row=row, column=4).value or "").strip()
        qty = parse_number(sheet.cell(row=row, column=5).value)

        if description == "" and first_cell and parse_number(first_cell) is None:
            section_text = normalize_text(first_cell)
            if section_text and not _is_non_product_description(section_text):
                current_section = str(first_cell).strip()
            continue

        if parse_number(first_cell) is None or not description or not unit or qty is None or qty <= 0:
            continue

        row_data = {
            "item_no": first_cell,
            "description": description,
            "unit": unit,
            "quantity": qty,
            "remark": current_section,
        }
        item = _build_takeoff_item(row_data, source_row=row)
        if item:
            item["parser_source"] = "YCBG A/B/D/E"
            parsed_items.append(item)

    parsed_items = assign_marks(parsed_items)
    parsed_items = apply_thickness_rules(parsed_items)
    return parsed_items

def _looks_like_order_takeoff_sheet(sheet) -> bool:
    for row in range(1, min(sheet.max_row, 25) + 1):
        text = " ".join(
            normalize_text(sheet.cell(row=row, column=col).value)
            for col in range(1, min(sheet.max_column, 20) + 1)
        )
        next_text = " ".join(
            normalize_text(sheet.cell(row=row + 1, column=col).value)
            for col in range(1, min(sheet.max_column, 20) + 1)
        )
        if (
            "qui cach ong gio" in text
            and "kich thuoc ong gio" in text
            and "so luong" in text
            and ("w mm" in next_text or "l mm" in next_text or "r" in next_text)
        ):
            return True
    return False


def _parse_order_takeoff_sheet(sheet) -> List[Dict[str, Any]]:
    parsed_items: List[Dict[str, Any]] = []
    for row in range(1, sheet.max_row + 1):
        item_no = sheet.cell(row=row, column=1).value
        if parse_number(item_no) is None:
            continue

        desc = str(sheet.cell(row=row, column=5).value or sheet.cell(row=row, column=4).value or "").strip()
        if not desc:
            continue

        kh_code = str(sheet.cell(row=row, column=3).value or "").strip()
        category = _category_from_order_kh(kh_code, desc)
        row_data = {
            "item_no": item_no,
            "mark": sheet.cell(row=row, column=2).value,
            "code": kh_code,
            "category": category,
            "description": desc,
            "width": sheet.cell(row=row, column=6).value,
            "height": sheet.cell(row=row, column=7).value,
            "width2": sheet.cell(row=row, column=8).value,
            "height2": sheet.cell(row=row, column=9).value,
            "length": sheet.cell(row=row, column=12).value,
            "radius": sheet.cell(row=row, column=13).value,
            "angle": sheet.cell(row=row, column=14).value,
            "unit": sheet.cell(row=row, column=15).value,
            "quantity": sheet.cell(row=row, column=16).value,
            "material": sheet.cell(row=row, column=17).value,
            "remark": sheet.cell(row=row, column=18).value,
        }
        item = _build_takeoff_item(row_data, source_row=row)
        if item:
            item["warnings"].append(
                "Dòng đọc từ mẫu đơn đặt hàng ống gió nhiều tầng tiêu đề; cần kiểm tra lại W/H/L/R/góc trước khi phát hành."
            )
            parsed_items.append(item)

    parsed_items = assign_marks(parsed_items)
    parsed_items = apply_thickness_rules(parsed_items)
    return parsed_items


def _category_from_order_kh(kh_code: Any, description: str) -> str:
    code = normalize_text(kh_code)
    if code in {"t", "ta"}:
        return "SUPPLY_DUCT"
    if code == "g":
        return "REDUCER"
    if code in {"cv", "c", "co"}:
        return "ELBOW"
    if code == "n":
        return "TRANSITION"
    return classify_product_category(description)


def _detail_quantity_total(sheet, header_row: int, row: int, quantity_col: int) -> float:
    header_text = " ".join(
        normalize_text(sheet.cell(header_row, col).value)
        for col in range(quantity_col + 1, min(sheet.max_column, quantity_col + 8) + 1)
    )
    subheader_text = " ".join(
        normalize_text(sheet.cell(header_row + 1, col).value)
        for col in range(quantity_col + 1, min(sheet.max_column, quantity_col + 8) + 1)
    )
    if not any(token in f"{header_text} {subheader_text}" for token in ["chi tiet", "block", "phan ham", "tttm"]):
        return 0.0

    total = 0.0
    for col in range(quantity_col + 1, min(sheet.max_column, quantity_col + 7) + 1):
        value = parse_number(sheet.cell(row=row, column=col).value)
        if value:
            total += value
    return total


def load_legacy_xls_takeoff(file_path: str):
    """
    Read old binary .xls files. Some Vietnamese estimating workbooks contain
    legacy Excel name formulas that make xlrd raise AssertionError while parsing
    globals. The data sheets are still readable, so we bypass name formula
    evaluation and copy the KL sheet into an openpyxl workbook.
    """
    original_evaluator = getattr(xlrd.book, "evaluate_name_formula", None)
    xlrd.book.evaluate_name_formula = lambda *args, **kwargs: None
    try:
        book = xlrd.open_workbook(file_path)
    finally:
        if original_evaluator is not None:
            xlrd.book.evaluate_name_formula = original_evaluator

    sheet_name = "KL" if "KL" in book.sheet_names() else book.sheet_names()[0]
    legacy_sheet = book.sheet_by_name(sheet_name)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    for row_idx in range(legacy_sheet.nrows):
        for col_idx in range(legacy_sheet.ncols):
            value = legacy_sheet.cell_value(row_idx, col_idx)
            if value not in ("", None):
                sheet.cell(row=row_idx + 1, column=col_idx + 1, value=value)
    return workbook, sheet


def extract_dimensions_from_text(text: str) -> Dict[str, float]:
    if not text:
        return {}

    normalized = normalize_text(text)
    dims: Dict[str, float] = {}
    number = r"(\d+(?:[\.,]\d+)?)"

    kt_3d = re.search(rf"(?:\bKT\b|\bkich thuoc\b)\s*{number}\s*[xX*×]\s*{number}\s*[xX*×]\s*{number}", text, re.IGNORECASE)
    kt_2d = re.search(rf"(?:\bKT\b|\bkich thuoc\b)\s*{number}\s*[xX*×]\s*{number}", text, re.IGNORECASE)

    if kt_3d:
        dims["width"] = parse_number(kt_3d.group(1))
        dims["height"] = parse_number(kt_3d.group(2))
        dims["length"] = parse_number(kt_3d.group(3))
    elif kt_2d:
        dims["width"] = parse_number(kt_2d.group(1))
        dims["height"] = parse_number(kt_2d.group(2))

    rect_3d = None if dims else re.search(rf"\b{number}\s*[xX*×]\s*{number}\s*[xX*×]\s*{number}\b", text)
    if rect_3d:
        dims["width"] = parse_number(rect_3d.group(1))
        dims["height"] = parse_number(rect_3d.group(2))
        dims["length"] = parse_number(rect_3d.group(3))
    else:
        rect_2d = None if dims else re.search(rf"\b{number}\s*[xX*×]\s*{number}\b", text)
        if rect_2d:
            dims["width"] = parse_number(rect_2d.group(1))
            dims["height"] = parse_number(rect_2d.group(2))

    round_match = re.search(rf"(?:\b[dD]|[Øø]|dia|phi)\s*{number}", text)
    if round_match:
        dims["diameter"] = parse_number(round_match.group(1))

    length_match = re.search(rf"(?:(?<![a-z])l|len|length|dai)\s*=?\s*{number}", normalized)
    if length_match:
        dims["length"] = parse_number(length_match.group(1))

    thickness_match = re.search(rf"(?:day|do day|ton day|dày|độ dày|tôn dày|chieu day|thick|tk|t\s*=)\s*{number}\s*(?:mm)?", normalized)
    if thickness_match:
        dims["thickness"] = parse_number(thickness_match.group(1))

    return {key: value for key, value in dims.items() if value is not None}


def extract_dimensions_from_text(text: str) -> Dict[str, float]:
    if not text:
        return {}

    normalized = normalize_text(text)
    dims: Dict[str, float] = {}
    number = r"(\d+(?:[\.,]\d+)?)"
    separator = r"[xX*×]"

    reducer_match = re.search(
        rf"{number}\s*{separator}\s*{number}\s*[-/]\s*{number}\s*{separator}\s*{number}\s*(?:[-/]\s*)?[lL]\s*{number}",
        text,
        re.IGNORECASE,
    )
    if reducer_match:
        dims["width"] = parse_number(reducer_match.group(1))
        dims["height"] = parse_number(reducer_match.group(2))
        dims["width2"] = parse_number(reducer_match.group(3))
        dims["height2"] = parse_number(reducer_match.group(4))
        dims["length"] = parse_number(reducer_match.group(5))

    transition_round_match = None if dims else re.search(
        rf"{number}\s*{separator}\s*{number}\s*[-/]\s*[ØøΦφ]?\s*{number}\s*[-/]\s*[lL]?\s*{number}",
        text,
        re.IGNORECASE,
    )
    if transition_round_match:
        dims["width"] = parse_number(transition_round_match.group(1))
        dims["height"] = parse_number(transition_round_match.group(2))
        dims["diameter"] = parse_number(transition_round_match.group(3))
        dims["length"] = parse_number(transition_round_match.group(4))

    kt_3d = None if dims else re.search(rf"(?:KT\s*:?\s*|\bkich thuoc\b\s*){number}\s*{separator}\s*{number}\s*{separator}\s*[lL]?\s*{number}", text, re.IGNORECASE)
    kt_2d = None if dims else re.search(rf"(?:KT\s*:?\s*|\bkich thuoc\b\s*){number}\s*{separator}\s*{number}", text, re.IGNORECASE)
    if kt_3d:
        dims["width"] = parse_number(kt_3d.group(1))
        dims["height"] = parse_number(kt_3d.group(2))
        dims["length"] = parse_number(kt_3d.group(3))
    elif kt_2d:
        dims["width"] = parse_number(kt_2d.group(1))
        dims["height"] = parse_number(kt_2d.group(2))

    rect_3d = None if dims else re.search(rf"\b{number}\s*{separator}\s*{number}\s*{separator}\s*[lL]?\s*{number}(?:\s*mm)?", text)
    rect_l = None if dims or rect_3d else re.search(rf"\b{number}\s*{separator}\s*{number}\s*[lL]\s*{number}(?:\s*mm)?", text, re.IGNORECASE)
    if rect_3d or rect_l:
        match = rect_3d or rect_l
        dims["width"] = parse_number(match.group(1))
        dims["height"] = parse_number(match.group(2))
        dims["length"] = parse_number(match.group(3))
    else:
        rect_2d = None if dims else re.search(rf"\b{number}\s*{separator}\s*{number}(?:\s*mm)?", text)
        if rect_2d:
            dims["width"] = parse_number(rect_2d.group(1))
            dims["height"] = parse_number(rect_2d.group(2))

    round_match = re.search(rf"(?:\b[dD]|[ØøΦφ]|dia|phi)\s*{number}", text, re.IGNORECASE)
    if round_match:
        dims["diameter"] = parse_number(round_match.group(1))

    round_length_match = re.search(rf"[ØøΦφ]?\s*{number}\s*{separator}\s*[lL]\s*{number}(?:\s*mm)?", text, re.IGNORECASE)
    if round_length_match and ("Ø" in text or "ø" in text or "phi" in normalized):
        dims["diameter"] = dims.get("diameter") or parse_number(round_length_match.group(1))
        dims["length"] = dims.get("length") or parse_number(round_length_match.group(2))

    length_match = re.search(rf"(?:(?<![a-z])l|len|length|dai)\s*=?\s*{number}", normalized)
    if length_match:
        dims["length"] = parse_number(length_match.group(1))

    thickness_match = re.search(rf"(?:day|do day|ton day|dày|độ dày|ton dày|tôn dày|chieu day|thick|tk|t\s*=)\s*{number}\s*(?:mm)?", normalized)
    if thickness_match:
        dims["thickness"] = parse_number(thickness_match.group(1))

    angle_match = re.search(r"(?<![a-z])(\d+(?:[\.,]\d+)?)\s*(?:°|do\b)(?!\s*day)", normalized)
    if angle_match:
        dims["angle"] = parse_number(angle_match.group(1))

    radius_match = re.search(rf"[rR]\s*{number}(?:\s*mm)?", text)
    if radius_match:
        dims["radius"] = parse_number(radius_match.group(1))

    return {key: value for key, value in dims.items() if value is not None}


def infer_material(text: str, fallback: str = "GI") -> str:
    normalized = normalize_text(text)
    if any(keyword in normalized for keyword in ["inox", "stainless", "ss"]):
        return "SS"
    if any(keyword in normalized for keyword in ["nhom", "aluminium", "aluminum"]):
        return "AL"
    if any(keyword in normalized for keyword in ["thep den", "mild steel", "ms"]):
        return "MS"
    if any(keyword in normalized for keyword in ["ton", "ma kem", "galvanized", "z080", "z275"]):
        return "GI"
    return fallback or "GI"


def infer_pressure_class(text: str) -> str:
    normalized = normalize_text(text)
    match = re.search(r"\bei\s*\d+", normalized)
    return match.group(0).upper().replace(" ", "") if match else ""


def classify_product_category(description: str, name: str = "") -> str:
    text = normalize_text(f"{name} {description}")
    try:
        from backend.product_catalog import match_product

        product_match = match_product(text)
        if product_match and product_match.get("category"):
            return str(product_match["category"])
    except Exception as exc:
        print(f"Product alias classification skipped: {exc}")

    direct_rules = [
        ("END_CAP", ["ong bit dau", "dau bit", "bit dau", "bit dau"]),
        ("PLENUM_BOX", ["hop gio", "hop plenum", "plenum"]),
        ("FLEXIBLE_DUCT", ["flexible air duct", "ong gio mem", "ong mem"]),
        ("INSULATION", ["ductwork insulation", "cach nhiet ong gio", "bao on ong gio", "cach nhiet"]),
        ("FILTER", ["g4 filter", "loc g4", "filter"]),
        ("LOUVER", ["ventcap", "vent cap", "nap thong gio"]),
        ("ACCESSORY", ["insect screen", "lcct", "luoi chan con trung", "luoi chong con trung"]),
        ("SQUARE_DIFFUSER", ["cua nan", "cua gio nan", "mieng gio", "cua gio", "air grill", "air grille", "eag", "teag", "grille", "diffuser"]),
        ("FLEXIBLE_CONNECTOR", ["khop noi mem", "co bat", "vai bat canvas", "simili"]),
        ("MOTORIZED_DAMPER", ["van mfd", "mfd", "motorized fire damper", "van chan lua dong mo bang mo to", "van chan lua dong mo bang motor"]),
        ("FIRE_DAMPER", ["fire damper", "van fd", "van chan lua", "fd"]),
        ("MOTORIZED_DAMPER", ["van md", "md", "van dieu chinh bang dong co"]),
        ("VOLUME_CONTROL_DAMPER", ["prd", "van giam ap", "van xa ap", "pressure relief damper"]),
        ("VOLUME_CONTROL_DAMPER", ["vcd", "van chinh luu", "volume control", "obd"]),
        ("BACK_DRAFT_DAMPER", ["nrd", "back draft", "van gio 1 chieu", "van 1 chieu", "van mot chieu"]),
        ("TRANSITION", ["vuong tron", "vuong/tron", "tron vuong", "con chuyen", "con dau quat", "dau quat"]),
        ("REDUCER", ["con thu", "con giam", "giam cap", "reducer"]),
        ("ELBOW", ["cut", "cut 45", "cut 90", "cút", "chech", "chech 45", "co 45", "co 90", "elbow"]),
        ("TRANSITION", ["got giay", "noi chan", "chan re"]),
        ("TEE", ["tee", "te", "chac 3", "ba nga"]),
        ("CROSS", ["cross", "tu nga"]),
        ("PLENUM_BOX", ["box mieng gio"]),
        ("LINEAR_DIFFUSER", ["linear", "slot diffuser"]),
        ("LOUVER", ["louver", "lover", "oal", "fal", "eal", "cua gio tuoi", "cua lay gio", "cua chop"]),
        ("EXHAUST_AIR_DUCT", ["ong gio thai", "hut thai", "exhaust"]),
        ("FRESH_AIR_DUCT", ["ong gio tuoi", "fresh air"]),
        ("SUPPLY_DUCT", ["ong cap gio", "cap gio", "supply duct"]),
        ("RETURN_DUCT", ["ong hoi gio", "hoi gio", "return duct"]),
        ("SMOKE_DUCT", ["hut khoi", "thai khoi", "smoke"]),
        ("SUPPLY_DUCT", ["ong", "ống"]),
        ("SUPPLY_DUCT", ["ong thong gio", "ong gio", "duct"]),
    ]
    for category, keywords in direct_rules:
        if any(_contains_keyword(text, keyword) for keyword in keywords):
            return category

    if supabase_store.enabled():
        try:
            categories = supabase_store.get_category_keywords()
        except Exception as exc:
            if _supabase_strict_enabled():
                raise
            print(f"Supabase get_category_keywords unavailable; using local dev DB because USE_SUPABASE_DB=0: {exc}")
            categories = []
    else:
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT normalized_key, keywords FROM product_categories")
        categories = cursor.fetchall()
        conn.close()

    for row in categories:
        cat_key = row["normalized_key"] if isinstance(row, dict) else row[0]
        keywords_value = row["keywords"] if isinstance(row, dict) else row[1]
        if cat_key == "UNKNOWN":
            continue
        keyword_parts = keywords_value if isinstance(keywords_value, list) else str(keywords_value or "").split(",")
        keywords = [normalize_text(kw) for kw in keyword_parts if str(kw).strip()]
        if any(keyword and _contains_keyword(text, keyword) for keyword in keywords):
            return cat_key

    return "UNKNOWN"


def _contains_keyword(text: str, keyword: str) -> bool:
    keyword = normalize_text(keyword)
    if not keyword:
        return False
    if " " not in keyword and keyword.isalnum():
        return re.search(rf"(?<![a-z0-9]){re.escape(keyword)}(?![a-z0-9])", text) is not None
    return keyword in text


def classify_and_extract_llm(description: str) -> Dict[str, Any]:
    try:
        from backend import llm_provider
    except Exception:
        return {}
    if not llm_provider.enabled():
        return {}

    try:
        prompt = f"""
Bạn là kỹ sư QS HVAC. Hãy đọc mô tả vật tư bằng tiếng Việt, tiếng Anh hoặc tiếng Trung,
chuẩn hóa tên vật tư sang tiếng Việt và trích xuất danh mục, kích thước, vật liệu.

Danh mục hợp lệ:
SUPPLY_DUCT, RETURN_DUCT, FRESH_AIR_DUCT, EXHAUST_AIR_DUCT, SMOKE_DUCT, FIRE_DAMPER, MOTORIZED_DAMPER, VOLUME_CONTROL_DAMPER, BACK_DRAFT_DAMPER, FLEXIBLE_DUCT, FLEXIBLE_CONNECTOR, REDUCER, ELBOW, TRANSITION, TEE, CROSS, OFFSET, ACCESS_DOOR, INSPECTION_DOOR, SILENCER, PLENUM_BOX, LINEAR_DIFFUSER, SQUARE_DIFFUSER, ROUND_DIFFUSER, EGGCRATE_GRILLE, JET_NOZZLE, LOUVER, UNKNOWN.

Mô tả gốc: "{description}"

Chỉ trả về JSON:
{{
  "description_vi": "tên vật tư tiếng Việt nếu có thể chuẩn hóa, nếu không thì giữ nguyên",
  "category": "category_key",
  "width": null,
  "height": null,
  "length": null,
  "diameter": null,
  "thickness": null,
  "material": null
}}
"""
        return llm_provider.generate_json(prompt, temperature=0.1)
    except Exception as exc:
        print(f"LLM classification failed: {exc}")
        return {}


def _extract_json_payload(text: str) -> Any:
    cleaned = (text or "").strip()
    cleaned = re.sub(r"^```(?:json)?", "", cleaned, flags=re.IGNORECASE).strip()
    cleaned = re.sub(r"```$", "", cleaned).strip()
    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        match = re.search(r"(\[[\s\S]*\]|\{[\s\S]*\})", cleaned)
        if not match:
            raise
        return json.loads(match.group(1))


def _build_takeoff_item(row_data: Dict[str, Any], source_row: int | None = None) -> Optional[Dict[str, Any]]:
    quantity = parse_number(row_data.get("quantity")) or 0.0
    unit = str(row_data.get("unit", "") or "").strip()
    desc = str(row_data.get("description", "") or "").strip()
    material_spec = str(row_data.get("material", "") or "").strip()
    remark = str(row_data.get("remark", "") or "").strip()

    if quantity <= 0 and not unit:
        return None
    if not desc:
        return None
    normalized_desc = normalize_text(desc)
    if _is_non_product_description(normalized_desc):
        return None
    if normalized_desc in ["ten vat tu", "ten vat tu, thiet bi", "noi dung", "description"]:
        return None
    if "noi dung cong viec" in normalized_desc or "工作名称" in desc:
        return None
    if normalize_text(row_data.get("item_no")) in ["stt", "tt", "item"]:
        return None

    full_text = " ".join(part for part in [desc, material_spec, remark] if part)
    product_match = None
    try:
        from backend.product_catalog import match_product

        product_match = match_product(full_text)
    except Exception as exc:
        if _supabase_strict_enabled():
            raise
        print(f"Product catalog match skipped: {exc}")
    dims = extract_dimensions_from_text(full_text)

    width = row_data.get("width") or dims.get("width")
    height = row_data.get("height") or dims.get("height")
    width2 = row_data.get("width2") or dims.get("width2")
    height2 = row_data.get("height2") or dims.get("height2")
    width3 = row_data.get("width3") or dims.get("width3")
    height3 = row_data.get("height3") or dims.get("height3")
    diameter = row_data.get("diameter") or dims.get("diameter")
    default_length = 1000.0 if normalize_text(unit) in {"m", "met", "meter", "met dai"} else 1200.0
    length = row_data.get("length") or dims.get("length", default_length)
    thickness = row_data.get("thickness") or dims.get("thickness")
    radius = row_data.get("radius") or dims.get("radius")
    angle = row_data.get("angle") or dims.get("angle")
    material = infer_material(full_text, str(row_data.get("material_code") or "GI"))
    pressure_class = infer_pressure_class(full_text)
    desc_category = classify_product_category(desc)
    matched_category = str(product_match.get("category")) if product_match else ""
    if matched_category == "INSULATION" and desc_category not in {"", "UNKNOWN", "INSULATION"}:
        matched_category = desc_category
    category = str(row_data.get("category") or "").strip()
    if not category:
        category = desc_category if desc_category not in {"", "UNKNOWN"} else matched_category or classify_product_category(full_text)

    parse_llm_enabled = os.getenv("ENABLE_PARSE_LLM", "0" if os.getenv("PRODUCT_ONLY_MODE", "1").strip().lower() not in {"0", "false", "no", "off"} else "1")
    if category == "UNKNOWN" and parse_llm_enabled.strip().lower() in {"1", "true", "yes", "on"}:
        llm_res = classify_and_extract_llm(full_text)
        if llm_res:
            category = llm_res.get("category") or "UNKNOWN"
            if llm_res.get("description_vi"):
                desc = str(llm_res.get("description_vi") or desc).strip()
            width = width or llm_res.get("width")
            height = height or llm_res.get("height")
            length = length or llm_res.get("length") or 1200.0
            diameter = diameter or llm_res.get("diameter")
            thickness = thickness or llm_res.get("thickness")
            material = (llm_res.get("material") or material).upper()

    return {
        "item_no": str(row_data.get("item_no", "") or "").strip(),
        "code": str(row_data.get("code", "") or "").strip(),
        "mark": str(row_data.get("mark", "") or "").strip(),
        "description": desc,
        "name": desc.split(",")[0] if desc else "HVAC Item",
        "category": category,
        "product_code": product_match.get("product_code") if product_match else "",
        "product_name_vi": product_match.get("display_name_vi") if product_match else "",
        "matched_alias": product_match.get("matched_alias") if product_match else "",
        "product_required_fields": product_match.get("required_fields") if product_match else [],
        "product_pricing_method": product_match.get("pricing_method") if product_match else "",
        "width": parse_number(width),
        "height": parse_number(height),
        "width2": parse_number(width2),
        "height2": parse_number(height2),
        "width3": parse_number(width3),
        "height3": parse_number(height3),
        "diameter": parse_number(diameter),
        "length": parse_number(length),
        "radius": parse_number(radius),
        "angle": parse_number(angle),
        "thickness": parse_number(thickness),
        "material": material,
        "source_material_spec": material_spec,
        "quote_material_spec": material_spec,
        "material_spec_source": "file" if material_spec else "unresolved",
        "pressure_class": pressure_class,
        "quantity": quantity,
        "unit": unit or "pcs",
        "remark": remark,
        "source_row": source_row,
        "warnings": [],
    }


def _is_non_product_description(normalized_desc: str) -> bool:
    if not normalized_desc:
        return True
    skip_exact = {
        "tong cong",
        "tong cong gia truoc thue",
        "tong cong gia sau thue",
        "tong thanh toan",
        "thue gtgt 10%",
        "thue gtgt 8%",
        "vat",
        "van chuyen",
        "viet bang chu",
        "ghi chu",
    }
    if normalized_desc in skip_exact:
        return True
    skip_prefixes = (
        "tong cong",
        "thue gtgt",
        "thue vat",
        "van chuyen",
        "bang chu",
        "ghi chu",
        "ben ban",
        "ben mua",
        "kinh gui",
        "dia chi",
        "so dt",
        "ma so thue",
        "n/v bao gia",
    )
    return normalized_desc.startswith(skip_prefixes)


def parse_image_takeoff(file_path: str) -> List[Dict[str, Any]]:
    mime_type = mimetypes.guess_type(file_path)[0] or "image/png"
    with open(file_path, "rb") as image_file:
        image_bytes = image_file.read()

    prompt = """
Bạn là kỹ sư QS HVAC. Hãy OCR ảnh/bản scan bảng khối lượng và trích xuất các dòng vật tư cần báo giá.
Ảnh có thể chứa tiếng Việt, tiếng Anh, tiếng Trung hoặc song ngữ. Hãy đọc được cả ba ngôn ngữ này.

Yêu cầu:
- Chỉ lấy dòng sản phẩm/vật tư, bỏ tiêu đề, tổng cộng, ghi chú, dòng trống.
- Cột description phải là tiếng Việt để đưa vào báo giá. Nếu ảnh là tiếng Anh/Trung, hãy dịch/chuẩn hóa sang tiếng Việt HVAC tự nhiên.
- Có thể giữ ký hiệu/kích thước gốc trong description, nhưng không để toàn bộ tên vật tư bằng tiếng Trung nếu đã hiểu nghĩa.
- Nếu thấy kích thước như KT 1100x600, ống 500x700xL1120mm, Ø660, dày 0.75mm thì điền cả vào mô tả và các trường width/height/length/diameter/thickness nếu chắc chắn.
- Nếu không chắc trường nào thì để null, không tự bịa.
- Trả về JSON array thuần, không markdown.

Schema mỗi dòng:
{
  "item_no": "STT nếu có",
  "code": "mã nếu có",
  "mark": "ký hiệu nếu có",
  "description": "tên vật tư/công tác bằng tiếng Việt",
  "unit": "đơn vị",
  "quantity": số_lượng,
  "material": "vật liệu/quy cách nếu có",
  "remark": "ghi chú nếu có",
  "width": null,
  "height": null,
  "length": null,
  "diameter": null,
  "thickness": null,
  "category": "UNKNOWN"
}
"""
    response = _generate_gemini_content(
        [
            prompt,
            {"mime_type": mime_type, "data": image_bytes},
        ],
        generation_config={"response_mime_type": "application/json"},
    )
    payload = _extract_json_payload(response.text)
    rows = payload.get("items", payload) if isinstance(payload, dict) else payload
    if not isinstance(rows, list):
        raise ValueError("AI OCR không trả về danh sách dòng khối lượng hợp lệ.")

    parsed_items = []
    for idx, row in enumerate(rows, start=1):
        if not isinstance(row, dict):
            continue
        item = _build_takeoff_item(row, source_row=idx)
        if item:
            item["warnings"].append("Dòng đọc từ ảnh OCR, cần kiểm tra lại trước khi xuất báo giá chính thức.")
            parsed_items.append(item)

    parsed_items = assign_marks(parsed_items)
    parsed_items = apply_thickness_rules(parsed_items)
    return parsed_items


def parse_customer_sheet(file_path: str) -> List[Dict[str, Any]]:
    if os.path.splitext(file_path.lower())[1] in IMAGE_EXTENSIONS:
        return parse_image_takeoff(file_path)

    if file_path.lower().endswith(".xls") and not file_path.lower().endswith(".xlsx"):
        workbook, sheet = load_legacy_xls_takeoff(file_path)
    else:
        workbook = openpyxl.load_workbook(file_path, data_only=True)
        sheet = select_takeoff_sheet(workbook)

    if _looks_like_no_header_quantity_matrix(sheet):
        parsed_items = _parse_no_header_quantity_matrix(sheet)
        workbook.close()
        return parsed_items

    if _looks_like_ycbg_request_sheet(sheet):
        parsed_items = _parse_ycbg_request_sheet(sheet)
        workbook.close()
        return parsed_items

    if _looks_like_order_takeoff_sheet(sheet):
        parsed_items = _parse_order_takeoff_sheet(sheet)
        workbook.close()
        return parsed_items

    merged_lookup = resolve_merged_cells(sheet)
    header_row, col_mapping = detect_headers(sheet, merged_lookup)

    parsed_items = []
    for row in range(header_row + 1, sheet.max_row + 1):
        row_data = {}
        has_value = False
        for field, col_idx in col_mapping.items():
            value = merged_lookup.get((row, col_idx), sheet.cell(row=row, column=col_idx).value)
            if value is not None and str(value).strip() != "":
                has_value = True
            row_data[field] = value
        quantity_col = col_mapping.get("quantity")
        if quantity_col and not parse_number(row_data.get("quantity")):
            detail_total = _detail_quantity_total(sheet, header_row, row, quantity_col)
            if detail_total > 0:
                row_data["quantity"] = detail_total
        if not has_value:
            continue

        item = _build_takeoff_item(row_data, source_row=row)
        if item and float(item.get("quantity") or 0) > 0:
            parsed_items.append(item)

    workbook.close()
    parsed_items = assign_marks(parsed_items)
    parsed_items = apply_thickness_rules(parsed_items)
    return parsed_items


def parse_reference_quote(file_path: str) -> List[Dict[str, Any]]:
    """
    Parse a completed quotation file as a price reference. For Kaiyo-style files,
    this reads the `báo giá` sheet and extracts description, qty, unit price and
    line total from the visible quotation columns.
    """
    if file_path.lower().endswith(".xls") and not file_path.lower().endswith(".xlsx"):
        parsed_xls = _parse_reference_quote_xls(file_path)
        return parsed_xls if parsed_xls else parse_customer_sheet(file_path)

    workbook = openpyxl.load_workbook(file_path, data_only=True)
    sheet, header_row = _select_reference_quote_sheet(workbook)

    if sheet is None or header_row is None:
        workbook.close()
        return parse_customer_sheet(file_path)

    reference_columns = _detect_reference_columns(sheet, header_row)
    formula_helpers = _collect_reference_formula_helpers(workbook)
    description_col = reference_columns.get("description", 2)
    quantity_col = reference_columns.get("quantity", 6)
    unit_price_col = reference_columns.get("unit_price", 7)
    line_total_col = reference_columns.get("line_total", 8)
    item_no_col = reference_columns.get("item_no", 1)
    code_col = reference_columns.get("code", 12)
    unit_col = reference_columns.get("unit", 5)
    material_col = reference_columns.get("material", 3)
    brand_col = reference_columns.get("brand", 4)
    remark_col = reference_columns.get("remark", 9)

    items: List[Dict[str, Any]] = []
    for row in range(header_row + 1, sheet.max_row + 1):
        desc = str(sheet.cell(row=row, column=description_col).value or "").strip()
        normalized_desc = normalize_text(desc)
        if _is_non_product_description(normalized_desc):
            continue
        qty = parse_number(sheet.cell(row=row, column=quantity_col).value) or 0.0
        unit_price = parse_number(sheet.cell(row=row, column=unit_price_col).value)
        line_total = parse_number(sheet.cell(row=row, column=line_total_col).value)
        if unit_price is None and line_total is not None and qty > 0:
            unit_price = line_total / qty
        if line_total is None and unit_price is not None and qty > 0:
            line_total = unit_price * qty
        if not desc or (unit_price is None and line_total is None):
            continue
        width = parse_number(sheet.cell(row=row, column=13).value)
        height = parse_number(sheet.cell(row=row, column=14).value)
        formula_width2 = parse_number(sheet.cell(row=row, column=15).value)
        formula_height2 = parse_number(sheet.cell(row=row, column=16).value)
        formula_width3 = parse_number(sheet.cell(row=row, column=17).value)
        formula_height3 = parse_number(sheet.cell(row=row, column=18).value)
        formula_length = parse_number(sheet.cell(row=row, column=19).value)
        formula_radius = parse_number(sheet.cell(row=row, column=20).value)
        formula_angle = parse_number(sheet.cell(row=row, column=21).value)
        formula_area = parse_number(sheet.cell(row=row, column=22).value)
        formula_unit_price = parse_number(sheet.cell(row=row, column=23).value)
        area_multiplier = parse_number(sheet.cell(row=row, column=24).value)
        accessory_area = parse_number(sheet.cell(row=row, column=25).value)
        accessory_unit_price = parse_number(sheet.cell(row=row, column=26).value)
        dims = extract_dimensions_from_text(desc)
        helper_formula = formula_helpers.get(normalized_desc, {})
        helper_code = str(helper_formula.get("quote_code") or "").strip()
        width = helper_formula.get("formula_width") or width or dims.get("width")
        height = helper_formula.get("formula_height") or height or dims.get("height")
        formula_width2 = helper_formula.get("formula_width2") if helper_formula.get("formula_width2") is not None else formula_width2
        formula_height2 = helper_formula.get("formula_height2") if helper_formula.get("formula_height2") is not None else formula_height2
        formula_width3 = helper_formula.get("formula_width3") if helper_formula.get("formula_width3") is not None else formula_width3
        formula_height3 = helper_formula.get("formula_height3") if helper_formula.get("formula_height3") is not None else formula_height3
        formula_length = helper_formula.get("formula_length") if helper_formula.get("formula_length") is not None else formula_length
        formula_radius = helper_formula.get("formula_radius") if helper_formula.get("formula_radius") is not None else formula_radius
        formula_angle = helper_formula.get("formula_angle") if helper_formula.get("formula_angle") is not None else formula_angle
        formula_area = helper_formula.get("formula_area") if helper_formula.get("formula_area") is not None else formula_area
        category = classify_product_category(desc)
        product_match = None
        try:
            from backend.product_catalog import match_product

            product_match = match_product(desc)
            if product_match and product_match.get("category"):
                category = str(product_match["category"])
        except Exception as exc:
            print(f"Reference product alias match skipped: {exc}")
        material_spec = str(sheet.cell(row=row, column=material_col).value or "").strip()

        items.append({
            "item_no": str(sheet.cell(row=row, column=item_no_col).value or "").strip(),
            "code": helper_code or str(sheet.cell(row=row, column=code_col).value or "").strip(),
            "mark": "",
            "description": desc,
            "name": desc,
            "category": category,
            "product_code": product_match.get("product_code") if product_match else "",
            "product_name_vi": product_match.get("display_name_vi") if product_match else "",
            "matched_alias": product_match.get("matched_alias") if product_match else "",
            "product_required_fields": product_match.get("required_fields") if product_match else [],
            "product_pricing_method": product_match.get("pricing_method") if product_match else "",
            "width": width,
            "height": height,
            "diameter": dims.get("diameter"),
            "length": formula_length or dims.get("length", 1200.0),
            "thickness": dims.get("thickness"),
            "material": infer_material(material_spec or desc, "GI"),
            "pressure_class": infer_pressure_class(material_spec or desc),
            "quantity": qty,
            "unit": str(sheet.cell(row=row, column=unit_col).value or "").strip(),
            "remark": str(sheet.cell(row=row, column=remark_col).value or "").strip(),
            "reference_unit_price": unit_price,
            "reference_line_total": line_total,
            "quote_material_spec": material_spec,
            "quote_brand": str(sheet.cell(row=row, column=brand_col).value or "").strip(),
            "quote_code": helper_code or str(sheet.cell(row=row, column=code_col).value or "").strip(),
            "quote_output_quantity": qty,
            "quote_output_unit": str(sheet.cell(row=row, column=unit_col).value or "").strip(),
            "quote_output_unit_price": unit_price,
            "formula_width": width,
            "formula_height": height,
            "formula_width2": formula_width2,
            "formula_height2": formula_height2,
            "formula_width3": formula_width3,
            "formula_height3": formula_height3,
            "formula_length": formula_length,
            "formula_radius": formula_radius,
            "formula_angle": formula_angle,
            "formula_area": formula_area,
            "formula_unit_price": formula_unit_price,
            "area_multiplier": area_multiplier,
            "accessory_area": accessory_area,
            "accessory_unit_price": accessory_unit_price,
            "warnings": [],
        })

    workbook.close()
    return items


def _collect_reference_formula_helpers(workbook) -> Dict[str, Dict[str, Any]]:
    helpers: Dict[str, Dict[str, Any]] = {}
    formula_aliases = {
        "quote_code": ["ma sp", "ma cong thuc", "code"],
        "formula_width": ["w1", "rong", "width"],
        "formula_height": ["h1", "cao", "height"],
        "formula_width2": ["w2"],
        "formula_height2": ["h2"],
        "formula_width3": ["w3"],
        "formula_height3": ["h3"],
        "formula_length": ["l/h", "l h", "dai", "length"],
        "formula_radius": ["r/d", "r d", "ban kinh", "duong kinh"],
        "formula_angle": [" e", "goc", "angle"],
        "formula_area": ["dien tich", "area"],
    }

    for sheet in workbook.worksheets:
        formula_header_row = None
        formula_cols: Dict[str, int] = {}
        for row in range(1, min(sheet.max_row, 80) + 1):
            row_text = " ".join(normalize_text(sheet.cell(row=row, column=col).value) for col in range(1, sheet.max_column + 1))
            if "ma sp" not in row_text or not any(token in row_text for token in ["w1", "h1", "dien tich"]):
                continue
            for col in range(1, sheet.max_column + 1):
                value = " " + normalize_text(sheet.cell(row=row, column=col).value) + " "
                compact_value = value.strip()
                for field, aliases in formula_aliases.items():
                    if field in formula_cols:
                        continue
                    if any(alias in value or compact_value == alias.strip() for alias in aliases):
                        formula_cols[field] = col
            if formula_cols.get("quote_code") and (formula_cols.get("formula_width") or formula_cols.get("formula_area")):
                formula_header_row = row
                break

        if not formula_header_row:
            continue

        product_header_row = None
        product_cols: Dict[str, int] = {}
        for row in range(formula_header_row, min(sheet.max_row, formula_header_row + 4) + 1):
            cols = _detect_reference_columns(sheet, row)
            if cols.get("description"):
                product_header_row = row
                product_cols = cols
                break
        if not product_header_row:
            continue

        description_col = product_cols.get("description")
        quantity_col = product_cols.get("quantity")
        unit_col = product_cols.get("unit")
        for row in range(product_header_row + 1, sheet.max_row + 1):
            desc = str(sheet.cell(row=row, column=description_col).value or "").strip()
            if not desc or _is_non_product_description(normalize_text(desc)):
                continue
            key = normalize_text(desc)
            helper: Dict[str, Any] = {}
            for field, col in formula_cols.items():
                value = sheet.cell(row=row, column=col).value
                if field == "quote_code":
                    helper[field] = str(value or "").strip()
                else:
                    helper[field] = parse_number(value)
            if quantity_col:
                helper["quantity"] = parse_number(sheet.cell(row=row, column=quantity_col).value)
            if unit_col:
                helper["unit"] = str(sheet.cell(row=row, column=unit_col).value or "").strip()
            if helper.get("quote_code") or any(helper.get(field) is not None for field in formula_cols if field != "quote_code"):
                helpers[key] = helper
    return helpers


def _parse_reference_quote_xls(file_path: str) -> List[Dict[str, Any]]:
    try:
        sheets = pd.read_excel(file_path, sheet_name=None, header=None, engine="xlrd")
    except Exception:
        return []

    best: Optional[Tuple[float, str, int, Dict[str, int]]] = None
    for sheet_name, frame in sheets.items():
        sheet_key = normalize_text(sheet_name)
        if sheet_key in {"kl", "boq", "khoi luong"} or "khoi luong" in sheet_key:
            continue
        header_row = _find_reference_header_row_frame(frame)
        if header_row is None:
            continue
        columns = _detect_reference_columns_frame(frame, header_row)
        if "description" not in columns or "quantity" not in columns:
            continue
        if "unit_price" not in columns and "line_total" not in columns:
            continue
        score = _score_reference_frame(frame, sheet_name, header_row, columns)
        if best is None or score > best[0]:
            best = (score, sheet_name, header_row, columns)

    if not best:
        return []

    _, sheet_name, header_row, columns = best
    frame = sheets[sheet_name]
    return _reference_items_from_frame(frame, header_row, columns)


def _frame_value(frame: pd.DataFrame, row: int, col: int) -> Any:
    if row < 0 or col < 0 or row >= len(frame.index) or col >= len(frame.columns):
        return None
    value = frame.iat[row, col]
    if pd.isna(value):
        return None
    return value


def _find_reference_header_row_frame(frame: pd.DataFrame) -> Optional[int]:
    max_rows = min(len(frame.index), 100)
    max_cols = min(len(frame.columns), 20)
    for row in range(max_rows):
        row_text = " ".join(normalize_text(_frame_value(frame, row, col)) for col in range(max_cols))
        has_description = any(keyword in row_text for keyword in ["ten vat tu", "mo ta chi tiet", "mo ta", "noi dung"])
        has_quantity = any(keyword in row_text for keyword in ["khoi luong", "so luong", "qty"])
        has_price = any(keyword in row_text for keyword in ["don gia", "thanh tien", "tong cong"])
        if has_description and has_quantity and has_price:
            return row
    return None


def _detect_reference_columns_frame(frame: pd.DataFrame, header_row: int) -> Dict[str, int]:
    mapping: Dict[str, int] = {}
    for col in range(len(frame.columns)):
        value = normalize_text(_frame_value(frame, header_row, col))
        next_value = normalize_text(_frame_value(frame, header_row + 1, col))
        merged_value = " ".join(part for part in [value, next_value] if part)
        if not merged_value:
            continue
        if value in ["stt", "tt"] or "stt" in value:
            mapping.setdefault("item_no", col)
        if "ma hieu" in merged_value or merged_value == "ma" or "code" in merged_value:
            mapping.setdefault("code", col)
        if "hang san xuat" in merged_value or "hang sx" in merged_value or "brand" in merged_value or "xuat xu" in merged_value:
            mapping.setdefault("brand", col)
        if "vat lieu che tao" in merged_value or "vat lieu" in merged_value or "material" in merged_value:
            mapping.setdefault("material", col)
        if (
            "ten vat tu" in merged_value
            or "san pham" in merged_value
            or "ten san pham" in merged_value
            or "mo ta chi tiet" in merged_value
            or merged_value == "mo ta"
            or "noi dung" in merged_value
            or "hang muc" in merged_value
        ):
            mapping.setdefault("description", col)
        if "don vi" in merged_value or "dvt" in merged_value:
            mapping.setdefault("unit", col)
        if "khoi luong" in merged_value or "so luong" in merged_value or "qty" in merged_value:
            mapping.setdefault("quantity", col)
        if "don gia" in value:
            mapping.setdefault("unit_price", col)
        if "thanh tien" in value or value == "tong cong":
            mapping.setdefault("line_total", col)
        if "ghi chu" in merged_value or "remark" in merged_value or "note" in merged_value:
            mapping.setdefault("remark", col)
    return mapping


def _score_reference_frame(frame: pd.DataFrame, sheet_name: str, header_row: int, columns: Dict[str, int]) -> float:
    description_col = columns.get("description", 1)
    quantity_col = columns.get("quantity", 5)
    unit_price_col = columns.get("unit_price")
    line_total_col = columns.get("line_total")
    product_rows = 0
    total_sum = 0.0
    for row in range(header_row + 1, len(frame.index)):
        desc = str(_frame_value(frame, row, description_col) or "").strip()
        if not desc or _is_non_product_description(normalize_text(desc)):
            continue
        qty = parse_number(_frame_value(frame, row, quantity_col)) or 0.0
        unit_price = parse_number(_frame_value(frame, row, unit_price_col)) if unit_price_col is not None else None
        line_total = parse_number(_frame_value(frame, row, line_total_col)) if line_total_col is not None else None
        if unit_price is None and line_total is not None and qty > 0:
            unit_price = line_total / qty
        if line_total is None and unit_price is not None and qty > 0:
            line_total = unit_price * qty
        if unit_price is None and line_total is None:
            continue
        product_rows += 1
        total_sum += float(line_total or 0)
    title_bonus = 1_000_000_000 if normalize_text(sheet_name) in {"bao gia", "quotation", "quote"} else 0
    return title_bonus + total_sum + product_rows * 1000


def _reference_items_from_frame(frame: pd.DataFrame, header_row: int, columns: Dict[str, int]) -> List[Dict[str, Any]]:
    description_col = columns.get("description", 1)
    quantity_col = columns.get("quantity", 5)
    unit_price_col = columns.get("unit_price")
    line_total_col = columns.get("line_total")
    item_no_col = columns.get("item_no", 0)
    code_col = columns.get("code", 11)
    unit_col = columns.get("unit", 4)
    material_col = columns.get("material", 2)
    brand_col = columns.get("brand", 3)
    remark_col = columns.get("remark", 8)
    items: List[Dict[str, Any]] = []

    for row in range(header_row + 1, len(frame.index)):
        desc = str(_frame_value(frame, row, description_col) or "").strip()
        normalized_desc = normalize_text(desc)
        if not desc or _is_non_product_description(normalized_desc):
            continue
        qty = parse_number(_frame_value(frame, row, quantity_col)) or 0.0
        unit_price = parse_number(_frame_value(frame, row, unit_price_col)) if unit_price_col is not None else None
        line_total = parse_number(_frame_value(frame, row, line_total_col)) if line_total_col is not None else None
        if unit_price is None and line_total is not None and qty > 0:
            unit_price = line_total / qty
        if line_total is None and unit_price is not None and qty > 0:
            line_total = unit_price * qty
        if unit_price is None and line_total is None:
            continue

        dims = extract_dimensions_from_text(desc)
        category = classify_product_category(desc)
        product_match = None
        try:
            from backend.product_catalog import match_product

            product_match = match_product(desc)
            if product_match and product_match.get("category"):
                category = str(product_match["category"])
        except Exception as exc:
            print(f"Reference XLS product alias match skipped: {exc}")
        material_spec = str(_frame_value(frame, row, material_col) or "").strip()
        unit = str(_frame_value(frame, row, unit_col) or "").strip()
        items.append({
            "item_no": str(_frame_value(frame, row, item_no_col) or "").strip(),
            "code": str(_frame_value(frame, row, code_col) or "").strip(),
            "mark": "",
            "description": desc,
            "name": desc,
            "category": category,
            "product_code": product_match.get("product_code") if product_match else "",
            "product_name_vi": product_match.get("display_name_vi") if product_match else "",
            "matched_alias": product_match.get("matched_alias") if product_match else "",
            "product_required_fields": product_match.get("required_fields") if product_match else [],
            "product_pricing_method": product_match.get("pricing_method") if product_match else "",
            "width": dims.get("width"),
            "height": dims.get("height"),
            "diameter": dims.get("diameter"),
            "length": formula_length or dims.get("length", 1200.0),
            "thickness": dims.get("thickness"),
            "material": infer_material(material_spec or desc, "GI"),
            "pressure_class": infer_pressure_class(material_spec or desc),
            "quantity": qty,
            "unit": unit,
            "remark": str(_frame_value(frame, row, remark_col) or "").strip(),
            "reference_unit_price": unit_price,
            "reference_line_total": line_total,
            "quote_material_spec": material_spec,
            "quote_brand": str(_frame_value(frame, row, brand_col) or "").strip(),
            "quote_code": str(_frame_value(frame, row, code_col) or "").strip(),
            "quote_output_quantity": qty,
            "quote_output_unit": unit,
            "quote_output_unit_price": unit_price,
            "formula_width": dims.get("width"),
            "formula_height": dims.get("height"),
            "formula_length": dims.get("length", 1200.0),
            "warnings": [],
        })
    return items


def _select_reference_quote_sheet(workbook) -> Tuple[Optional[Any], Optional[int]]:
    candidates = []
    for sheet in workbook.worksheets:
        header_row = _find_reference_header_row(sheet)
        if not header_row:
            continue
        sheet_key = normalize_text(sheet.title)
        if sheet_key in {"kl", "boq", "khoi luong"} or "khoi luong" in sheet_key:
            continue
        columns = _detect_reference_columns(sheet, header_row)
        description_col = columns.get("description", 2)
        quantity_col = columns.get("quantity", 6)
        unit_price_col = columns.get("unit_price")
        line_total_col = columns.get("line_total")
        if not unit_price_col and not line_total_col:
            continue

        product_rows = 0
        total_sum = 0.0
        unit_sum = 0.0
        for row in range(header_row + 1, sheet.max_row + 1):
            desc = str(sheet.cell(row=row, column=description_col).value or "").strip()
            if not desc or _is_non_product_description(normalize_text(desc)):
                continue
            qty = parse_number(sheet.cell(row=row, column=quantity_col).value) or 0.0
            unit_price = parse_number(sheet.cell(row=row, column=unit_price_col).value) if unit_price_col else None
            line_total = parse_number(sheet.cell(row=row, column=line_total_col).value) if line_total_col else None
            if unit_price is None and line_total is not None and qty > 0:
                unit_price = line_total / qty
            if line_total is None and unit_price is not None and qty > 0:
                line_total = unit_price * qty
            if unit_price is None and line_total is None:
                continue
            product_rows += 1
            unit_sum += float(unit_price or 0)
            total_sum += float(line_total or 0)

        title_bonus = 1_000_000_000 if sheet_key in {"bao gia", "quotation", "quote"} else 0
        score = title_bonus + total_sum + product_rows * 1000
        candidates.append((score, sheet, header_row))

    if not candidates:
        return None, None
    candidates.sort(key=lambda row: row[0], reverse=True)
    return candidates[0][1], candidates[0][2]


def _find_reference_header_row(sheet) -> Optional[int]:
    for row in range(1, min(sheet.max_row, 100) + 1):
        row_text = " ".join(
            normalize_text(sheet.cell(row=row, column=col).value)
            for col in range(1, min(sheet.max_column, 12) + 1)
        )
        has_description = any(keyword in row_text for keyword in ["ten vat tu", "mo ta chi tiet", "mo ta", "noi dung"])
        has_quantity = any(keyword in row_text for keyword in ["khoi luong", "so luong", "qty"])
        has_price = any(keyword in row_text for keyword in ["don gia", "thanh tien", "tong cong"])
        if has_description and has_quantity and has_price:
            return row
    return None


def _detect_reference_columns(sheet, header_row: int) -> Dict[str, int]:
    mapping: Dict[str, int] = {}
    for col in range(1, sheet.max_column + 1):
        value = normalize_text(sheet.cell(row=header_row, column=col).value)
        next_value = normalize_text(sheet.cell(row=header_row + 1, column=col).value)
        merged_value = " ".join(part for part in [value, next_value] if part)
        if not merged_value:
            continue

        if value in ["stt", "tt"] or "stt" in value:
            mapping.setdefault("item_no", col)
        if "ma hieu" in merged_value or merged_value == "ma" or "code" in merged_value:
            mapping.setdefault("code", col)
        if "hang san xuat" in merged_value or "hang sx" in merged_value or "brand" in merged_value:
            mapping.setdefault("brand", col)
        if "vat lieu che tao" in merged_value or "vat lieu" in merged_value or "material" in merged_value:
            mapping.setdefault("material", col)
        if (
            "ten vat tu" in merged_value
            or "san pham" in merged_value
            or "ten san pham" in merged_value
            or "mo ta chi tiet" in merged_value
            or merged_value == "mo ta"
            or "noi dung" in merged_value
            or "hang muc" in merged_value
        ):
            mapping.setdefault("description", col)
        if "don vi" in merged_value or "dvt" in merged_value:
            mapping.setdefault("unit", col)
        if "khoi luong" in merged_value or "so luong" in merged_value or "qty" in merged_value:
            mapping.setdefault("quantity", col)
        if "don gia" in value:
            mapping.setdefault("unit_price", col)
        if "thanh tien" in value or value == "tong cong":
            mapping.setdefault("line_total", col)
        if "ghi chu" in merged_value or "remark" in merged_value or "note" in merged_value:
            mapping.setdefault("remark", col)
    return mapping


def assign_marks(items: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    if supabase_store.enabled():
        try:
            prefixes = supabase_store.get_mark_prefixes()
        except Exception as exc:
            if _supabase_strict_enabled():
                raise
            print(f"Supabase get_mark_prefixes unavailable; using local dev DB because USE_SUPABASE_DB=0: {exc}")
            prefixes = {}
    else:
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT category, prefix FROM mark_rules")
        prefixes = {row[0]: row[1] for row in cursor.fetchall()}
        conn.close()

    existing_marks = set(item["mark"] for item in items if item["mark"])
    counters: Dict[str, int] = {}

    for item in items:
        if item["mark"]:
            continue
        prefix = prefixes.get(item["category"], "HVAC")
        idx = counters.get(prefix, 1)
        new_mark = f"{prefix}-{idx:03d}"
        while new_mark in existing_marks:
            idx += 1
            new_mark = f"{prefix}-{idx:03d}"
        item["mark"] = new_mark
        existing_marks.add(new_mark)
        counters[prefix] = idx + 1

    return items



def _material_tokens(value: Any) -> set[str]:
    return {token for token in normalize_text(str(value or "")).split() if len(token) > 1}


def _same_measurement(left: Any, right: Any, tolerance: float = 1.0) -> bool:
    left_number = parse_number(left)
    right_number = parse_number(right)
    return left_number is not None and right_number is not None and abs(left_number - right_number) <= tolerance


def _material_spec_is_specific(value: Any) -> bool:
    normalized = normalize_text(str(value or ""))
    if not normalized:
        return False
    return normalized not in {"gi", "ss", "ms", "al", "ton ma kem", "inox", "thep den", "nhom"}


def _material_memory_score(item: Dict[str, Any], candidate: Dict[str, Any]) -> float:
    if item.get("category") != candidate.get("category"):
        return -1.0

    score = 45.0
    if item.get("product_code") and item.get("product_code") == candidate.get("product_code"):
        score += 35.0
    if item.get("material") and item.get("material") == candidate.get("material"):
        score += 12.0
    if item.get("thickness") and _same_measurement(item.get("thickness"), candidate.get("thickness"), 0.02):
        score += 18.0

    for key in ("width", "height", "diameter", "length"):
        if item.get(key) and candidate.get(key) and _same_measurement(item.get(key), candidate.get(key)):
            score += 6.0

    pressure_class = str(item.get("pressure_class") or "")
    candidate_text = " ".join(str(candidate.get(key) or "") for key in ("quote_material_spec", "quote_note", "normalized_description"))
    if pressure_class and normalize_text(pressure_class) in normalize_text(candidate_text):
        score += 15.0

    item_tokens = _material_tokens(" ".join(str(item.get(key) or "") for key in ("description", "remark")))
    candidate_tokens = _material_tokens(candidate_text)
    if item_tokens and candidate_tokens:
        score += 20.0 * len(item_tokens & candidate_tokens) / max(1, len(item_tokens | candidate_tokens))
    return score


def resolve_material_specifications(
    items: List[Dict[str, Any]],
    memory_rows: Optional[List[Dict[str, Any]]] = None,
    product_rules: Optional[Dict[str, Dict[str, Any]]] = None,
) -> List[Dict[str, Any]]:
    """Resolve a full manufacturing specification without silently inventing one."""
    if memory_rows is None:
        memory_rows = []
        if supabase_store.enabled():
            try:
                memory_rows = supabase_store.get_approved_material_memory()
            except Exception as exc:
                if _supabase_strict_enabled():
                    raise
                print(f"Supabase material memory unavailable: {exc}")
    if product_rules is None:
        product_rules = {}
        if supabase_store.enabled():
            try:
                product_rules = supabase_store.get_product_pricing_rules()
            except Exception as exc:
                if _supabase_strict_enabled():
                    raise
                print(f"Supabase product rules unavailable: {exc}")

    for item in items:
        source_spec = str(item.get("source_material_spec") or item.get("quote_material_spec") or "").strip()
        if _material_spec_is_specific(source_spec):
            item["quote_material_spec"] = source_spec
            item["material_spec_source"] = "file"
            item["material_spec_confidence"] = 100
            continue

        ranked = sorted(
            ((_material_memory_score(item, candidate), candidate) for candidate in memory_rows
             if _material_spec_is_specific(candidate.get("quote_material_spec"))),
            key=lambda pair: pair[0],
            reverse=True,
        )
        best_score, best = ranked[0] if ranked else (-1.0, None)
        second_score, second = ranked[1] if len(ranked) > 1 else (-1.0, None)
        best_spec = str((best or {}).get("quote_material_spec") or "").strip()
        second_spec = str((second or {}).get("quote_material_spec") or "").strip()
        unambiguous = best_score >= 70 and (second_score < best_score - 6 or best_spec == second_spec)

        if unambiguous:
            item["quote_material_spec"] = best_spec
            item["material_spec_source"] = "supabase_memory"
            item["material_spec_confidence"] = round(min(100.0, best_score), 1)
            item["material_spec_reference"] = best.get("source_file", "")
            continue

        rule = product_rules.get(str(item.get("category") or ""), {})
        rule_spec = str(rule.get("material_spec") or "").strip()
        if _material_spec_is_specific(rule_spec):
            item["quote_material_spec"] = rule_spec
            item["material_spec_source"] = "approved_rule"
            item["material_spec_confidence"] = 95
            continue

        item["quote_material_spec"] = ""
        item["material_spec_source"] = "needs_confirmation"
        item["material_spec_confidence"] = 0
        item.setdefault("warnings", []).append(
            "Chưa xác nhận được vật liệu chế tạo đầy đủ (loại vật liệu/cấu tạo/độ dày). Sales/QS cần chọn mẫu vật liệu đã duyệt trước khi xuất."
        )
    return items

def apply_thickness_rules(items: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    if supabase_store.enabled():
        try:
            rules = supabase_store.get_thickness_rules()
        except Exception as exc:
            if _supabase_strict_enabled():
                raise
            print(f"Supabase get_thickness_rules unavailable; using local dev DB because USE_SUPABASE_DB=0: {exc}")
            rules = []
    else:
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT material, min_width, max_width, thickness FROM thickness_rules")
        rules = cursor.fetchall()
        conn.close()

    for item in items:
        if item.get("thickness"):
            continue

        material = item["material"]
        governing_size = max(item.get("width") or 0, item.get("height") or 0, item.get("diameter") or 0)
        matched_thickness = None
        for rule in rules:
            if isinstance(rule, dict):
                rule_mat = rule["material"]
                min_w = rule["min_width"]
                max_w = rule["max_width"]
                thickness = rule["thickness"]
            else:
                rule_mat, min_w, max_w, thickness = rule
            if rule_mat == material and min_w <= governing_size < max_w:
                matched_thickness = thickness
                break

        item["thickness"] = matched_thickness or (0.6 if material == "GI" else 0.8)

    return resolve_material_specifications(items)








