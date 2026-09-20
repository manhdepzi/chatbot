from io import BytesIO
import os

os.environ["PRODUCT_ONLY_MODE"] = "0"

from openpyxl import Workbook, load_workbook

from backend.excel_engine import build_quotation_workbook, _infer_template_quote_code


def test_kaiyo_template_mapping_smoke():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "báo giá"
    sheet["H13"] = "Dự án"
    sheet["A9"] = "Kính Gửi:"
    sheet["A21"] = "STT"
    sheet["B21"] = "Tên vật tư, thiết bị"
    sheet["C21"] = "Thông số kỹ thuật"
    sheet["D21"] = "Nhãn hiệu"
    sheet["E21"] = "Đơn vị"
    sheet["F21"] = "Khối lượng"
    sheet["G21"] = "Đơn giá"
    sheet["H21"] = "Thành tiền"
    sheet["I21"] = "Ghi chú"
    sheet["K20"] = "Tên vật tư, thiết bị"
    sheet["L21"] = "Mã SP"
    sheet["M21"] = "W1"
    sheet["N21"] = "H1"

    template = BytesIO()
    workbook.save(template)

    items = [
        {
            "item_no": "1",
            "description": "Co 700x400",
            "quote_material_spec": "Ton ma kem day 0.75 mm",
            "quote_brand": "Kaiyo Viet Nam",
            "unit": "cai",
            "quantity": 2,
            "quote_unit_price": 1147828,
            "quote_note": "",
            "quote_code": "cv",
            "width": 700,
            "height": 400,
            "calculated_area": 2.4,
            "pricing_mode": "area",
            "material": "GI",
            "thickness": 0.75,
        }
    ]
    summary = {"items_subtotal": 2295656, "vat": 229565.6, "grand_total": 2525221.6}

    output = build_quotation_workbook(
        items,
        summary,
        template_bytes=template.getvalue(),
        project="Project A",
        customer="Customer B",
    )
    result = load_workbook(BytesIO(output), data_only=False)
    quote = result["báo giá"]

    assert quote["I13"].value == "Project A"
    assert quote["C9"].value == "Customer B"
    assert quote["B24"].value == "Co 700x400"
    assert quote["C24"].value == "Ton ma kem day 0.75 mm"
    assert quote["D24"].value == "Kaiyo Viet Nam"
    assert quote["F24"].value == 2
    assert quote["G24"].value == 1147828
    assert quote["H24"].value == 2295656
    assert quote["L24"].value == "cv"
    assert quote["M24"].value == 700
    assert quote["N24"].value == 400


def test_kaiyo_template_mapping_does_not_require_sheet_name():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Mẫu - xây lắp"
    sheet["A21"] = "STT"
    sheet["B21"] = "Tên vật tư, thiết bị"
    sheet["C21"] = "Vật liệu chế tạo"
    sheet["D21"] = "Xuất Xứ"
    sheet["E21"] = "Đơn vị"
    sheet["F21"] = "Khối lượng"
    sheet["G21"] = "Đơn giá"
    sheet["H21"] = "Thành tiền"
    sheet["I21"] = "Ghi chú"

    template = BytesIO()
    workbook.save(template)

    output = build_quotation_workbook(
        [
            {
                "item_no": "1",
                "description": "Cửa gió nan thẳng",
                "quote_material_spec": "Nhôm định hình",
                "quote_brand": "Kaiyo Việt Nam",
                "unit": "cái",
                "quantity": 3,
                "quote_unit_price": 100000,
                "subtotal": 300000,
            }
        ],
        {"items_subtotal": 300000, "vat": 30000, "grand_total": 330000},
        template_bytes=template.getvalue(),
    )

    result = load_workbook(BytesIO(output), data_only=False)

    assert "Dữ liệu báo giá AI" not in result.sheetnames
    assert "AI Quotation Data" not in result.sheetnames
    quote = result["Mẫu - xây lắp"]
    assert quote["B24"].value == "Cửa gió nan thẳng"
    assert quote["F24"].value == 3
    assert quote["G24"].value == 100000
    assert quote["H24"].value == 300000


def test_kaiyo_template_keeps_area_formula_for_end_cap(monkeypatch):
    monkeypatch.setenv("PRODUCT_ONLY_MODE", "1")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "bao gia"
    sheet["A21"] = "STT"
    sheet["B21"] = "Ten vat tu, thiet bi"
    sheet["C21"] = "Vat lieu che tao"
    sheet["D21"] = "Xuat Xu"
    sheet["E21"] = "Don vi"
    sheet["F21"] = "Khoi luong"
    sheet["G21"] = "Don gia"
    sheet["H21"] = "Thanh tien"
    sheet["I21"] = "Ghi chu"
    sheet["K20"] = "Ten vat tu, thiet bi"
    sheet["L21"] = "Ma SP"
    sheet["M21"] = "W1"
    sheet["N21"] = "H1"
    sheet["S21"] = "L/H"
    sheet["V21"] = "DIEN TICH/Cai"

    template = BytesIO()
    workbook.save(template)

    output = build_quotation_workbook(
        [
            {
                "item_no": "1",
                "description": "Dau bit ton ma kem 1800x500L100mm",
                "category": "END_CAP",
                "quote_code": "tb",
                "formula_width": 1800,
                "formula_height": 500,
                "formula_length": 100,
                "formula_area": 0.46,
                "quantity": 2,
                "unit": "cai",
            }
        ],
        {"total_area": 2.72, "items_subtotal": 0, "vat": 0, "grand_total": 0},
        template_bytes=template.getvalue(),
    )

    result = load_workbook(BytesIO(output), data_only=False)
    quote = result["bao gia"]
    assert quote["L24"].value == "tb"
    assert quote["M24"].value == 1800
    assert quote["N24"].value == 500
    assert quote["S24"].value == 100
    assert isinstance(quote["V24"].value, str)
    assert quote["V24"].value.startswith("=IF(")
    assert 'L24="tb"' in quote["V24"].value
    assert "M24*N24" in quote["V24"].value


def test_infer_template_quote_code_from_product_code():
    assert _infer_template_quote_code({"product_code": "RECT_ELBOW"}) == "cv"
    assert _infer_template_quote_code({"product_code": "MFD_L250"}) == "mfd"
    assert _infer_template_quote_code({"product_code": "LOUVER_WITH_INSECT_SCREEN"}) == "c"


def test_infer_template_quote_code_from_category():
    assert _infer_template_quote_code({"category": "SUPPLY_DUCT"}) == "t"
    assert _infer_template_quote_code({"category": "VOLUME_CONTROL_DAMPER"}) == "vcd"
    assert _infer_template_quote_code({"category": "END_CAP"}) == "tb"


def test_infer_template_quote_code_transition_and_special_shapes():
    assert _infer_template_quote_code({"category": "TRANSITION"}) == "vt"
    assert _infer_template_quote_code({"category": "TRANSITION", "description": "Nối chân 400x300"}) == "n"
    assert _infer_template_quote_code({"category": "TRANSITION", "description": "Chân rẽ vuông"}) == "n"
    assert _infer_template_quote_code({"category": "ELBOW", "description": "Zét 90 độ"}) == "d"


def test_infer_template_quote_code_unknown_is_empty():
    assert _infer_template_quote_code({"category": "UNKNOWN"}) == ""
    assert _infer_template_quote_code({"category": "INSULATION"}) == ""
