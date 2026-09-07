import os

os.environ["USE_SUPABASE_DB"] = "0"

from openpyxl import Workbook

from backend.parser_engine import parse_customer_sheet, parse_reference_quote, resolve_material_specifications


def test_reference_quote_prefers_sales_quote_sheet_over_kl_sheet(tmp_path):
    workbook = Workbook()
    kl = workbook.active
    kl.title = "KL"
    kl["A6"] = "STT"
    kl["B6"] = "MÃ HIỆU\nĐƠN GIÁ"
    kl["C6"] = "NỘI DUNG CÔNG VIỆC"
    kl["D6"] = "ĐƠN VỊ"
    kl["E6"] = "KHỐI LƯỢNG"
    kl["A7"] = 1
    kl["B7"] = "BB.82404"
    kl["C7"] = "Louver 300x300"
    kl["D7"] = "cái"
    kl["E7"] = 2

    quote = workbook.create_sheet("Sheet1")
    quote["A21"] = "STT"
    quote["B21"] = "Tên vật tư, thiết bị"
    quote["E21"] = "Đơn vị"
    quote["F21"] = "Khối lượng"
    quote["G21"] = "Đơn giá"
    quote["H21"] = "Thành tiền"
    quote["A22"] = 1
    quote["B22"] = "Louver 300x300"
    quote["E22"] = "cái"
    quote["F22"] = 2
    quote["G22"] = 291000
    quote["H22"] = 582000

    path = tmp_path / "completed_quote.xlsx"
    workbook.save(path)

    rows = parse_reference_quote(str(path))

    assert len(rows) == 1
    assert rows[0]["reference_unit_price"] == 291000
    assert rows[0]["reference_line_total"] == 582000


def test_order_takeoff_sheet_with_two_header_rows(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet["E8"] = "KHỐI LƯỢNG ĐẶT HÀNG ĐƠN 1"
    headers = [
        "STT",
        "Ký hiệu",
        "KH",
        "QUI CÁCH ỐNG GIÓ",
        "QUI CÁCH ỐNG GIÓ",
        "KÍCH THƯỚC ỐNG GIÓ",
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        "Đơn vị",
        "Số lượng",
        "Thông số kỹ thuật ",
        "Ghi chú ",
    ]
    for col, value in enumerate(headers, start=1):
        sheet.cell(row=12, column=col, value=value)
    for col, value in {6: "W (mm)", 7: "H (mm)", 8: "W' (mm)", 9: "H' (mm)", 12: "L (mm)", 13: "R(Ø)\n(mm)", 14: "ĐỘ LỆCH /Góc Co(B)"}.items():
        sheet.cell(row=13, column=col, value=value)
    row = [
        1,
        "RA3-1",
        "t",
        "Ống CN",
        "Ống CN 1200x800xL1110",
        1200,
        800,
        None,
        None,
        None,
        None,
        1110,
        None,
        None,
        "cái",
        25,
        "Tôn mạ kẽm Z080 độ dày 0,75",
        None,
    ]
    for col, value in enumerate(row, start=1):
        sheet.cell(row=16, column=col, value=value)

    path = tmp_path / "order_takeoff.xlsx"
    workbook.save(path)

    rows = parse_customer_sheet(str(path))

    assert len(rows) == 1
    assert rows[0]["category"] == "SUPPLY_DUCT"
    assert rows[0]["width"] == 1200
    assert rows[0]["height"] == 800
    assert rows[0]["length"] == 1110
    assert rows[0]["quantity"] == 25
    assert rows[0]["thickness"] == 0.75
    assert rows[0].get("angle") is None



def test_fire_rated_duct_request_uses_product_column_for_dimensions(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    headers = ["STT", "S\u1ea3n ph\u1ea9m", "Quy c\u00e1ch ch\u1ed1ng ch\u00e1y", "\u0110VT", "S\u1ed1 l\u01b0\u1ee3ng", "Ghi ch\u00fa"]
    for col, value in enumerate(headers, start=1):
        sheet.cell(row=1, column=col, value=value)

    rows = [
        [1, "\u1ed0ng t\u00f4n m\u1ea1 k\u1ebdm 1800x500L1110mm, t\u00f4n d\u00e0y 0.75mm b\u1ecdc ch\u1ed1ng ch\u00e1y MGO EI45", "C\u1ea5u t\u1ea1o: T\u00f4n Hoa Sen m\u1ea1 k\u1ebdm d\u00e0y 0.75mm v\u00e0 t\u1ea5m MGO ng\u0103n ch\u00e1y c\u00e1ch nhi\u1ec7t Feco MF60 d\u00e0y 5mm", "\u0110o\u1ea1n", 97, "\u0110\u00e3 bao g\u1ed3m t\u0103ng c\u1ee9ng"],
        [2, "\u0110\u1ea7u b\u1ecbt t\u00f4n m\u1ea1 k\u1ebdm 1500x400L100mm, t\u00f4n d\u00e0y 0.75mm b\u1ecdc ch\u1ed1ng ch\u00e1y MGO EI45", "C\u1ea5u t\u1ea1o: T\u00f4n Hoa Sen m\u1ea1 k\u1ebdm d\u00e0y 0.75mm v\u00e0 t\u1ea5m MGO ng\u0103n ch\u00e1y c\u00e1ch nhi\u1ec7t Feco MF120 d\u00e0y 20mm", "\u0110o\u1ea1n", 7, ""],
        [3, "Ch\u1ebfch 45 \u0111\u1ed9 t\u00f4n m\u1ea1 k\u1ebdm 0.75mm b\u1ecdc EI45 KT 1800x500 R500", "C\u1ea5u t\u1ea1o: T\u00f4n Hoa Sen m\u1ea1 k\u1ebdm d\u00e0y 0.75mm v\u00e0 t\u1ea5m MGO ng\u0103n ch\u00e1y c\u00e1ch nhi\u1ec7t Feco MF60 d\u00e0y 5mm", "c\u00e1i", 2, ""],
        [4, "C\u00f4n thu t\u00f4n m\u1ea1 k\u1ebdm 0.75mm b\u1ecdc EI45 KT 1800x500/1500x400L500 \u0111\u1ec1u, ph\u1eb3ng \u0111\u00e1y", "C\u1ea5u t\u1ea1o: T\u00f4n Hoa Sen m\u1ea1 k\u1ebdm d\u00e0y 0.75mm v\u00e0 t\u1ea5m MGO ng\u0103n ch\u00e1y c\u00e1ch nhi\u1ec7t Feco MF60 d\u00e0y 5mm", "c\u00e1i", 3, ""],
    ]
    for row_idx, row in enumerate(rows, start=2):
        for col, value in enumerate(row, start=1):
            sheet.cell(row=row_idx, column=col, value=value)

    path = tmp_path / "fire_rated_duct_request.xlsx"
    workbook.save(path)

    parsed = parse_customer_sheet(str(path))

    assert len(parsed) == 4
    assert parsed[0]["description"].startswith("\u1ed0ng t\u00f4n")
    assert parsed[0]["category"] == "SUPPLY_DUCT"
    assert parsed[0]["width"] == 1800
    assert parsed[0]["height"] == 500
    assert parsed[0]["length"] == 1110
    assert parsed[1]["category"] == "END_CAP"
    assert parsed[1]["length"] == 100
    assert parsed[2]["category"] == "ELBOW"
    assert parsed[2]["angle"] == 45
    assert parsed[2]["radius"] == 500
    assert parsed[3]["category"] == "REDUCER"
    assert parsed[3]["width2"] == 1500
    assert parsed[3]["height2"] == 400
    assert parsed[3]["length"] == 500


def test_material_spec_prefers_explicit_source_file_specification():
    item = {
        "category": "SMOKE_DUCT",
        "description": "Ong gio hut khoi 1800x500",
        "source_material_spec": "Cau tao: Ton Hoa Sen ma kem day 0.75mm va tam MGO EI45 day 5mm",
        "material": "GI",
        "thickness": 0.75,
        "warnings": [],
    }

    result = resolve_material_specifications([item], memory_rows=[], product_rules={})[0]

    assert result["quote_material_spec"] == item["source_material_spec"]
    assert result["material_spec_source"] == "file"
    assert result["material_spec_confidence"] == 100


def test_material_spec_uses_unambiguous_approved_supabase_memory():
    item = {
        "category": "SMOKE_DUCT",
        "product_code": "SMOKE_DUCT",
        "description": "Ong gio hut khoi EI45 1800x500 L1000",
        "remark": "",
        "material": "GI",
        "pressure_class": "EI45",
        "width": 1800,
        "height": 500,
        "length": 1000,
        "thickness": 0.75,
        "warnings": [],
    }
    memory = [{
        "category": "SMOKE_DUCT",
        "product_code": "SMOKE_DUCT",
        "normalized_description": "ong gio hut khoi ei45 1800x500 l1000",
        "material": "GI",
        "width": 1800,
        "height": 500,
        "length": 1000,
        "thickness": 0.75,
        "quote_material_spec": "Ton Hoa Sen ma kem day 0.75mm boc MGO EI45 day 5mm",
        "quote_note": "",
        "source_file": "approved-smoke-duct.xlsx",
    }]

    result = resolve_material_specifications([item], memory_rows=memory, product_rules={})[0]

    assert result["quote_material_spec"] == memory[0]["quote_material_spec"]
    assert result["material_spec_source"] == "supabase_memory"
    assert result["material_spec_reference"] == "approved-smoke-duct.xlsx"


def test_material_spec_warns_when_approved_memory_is_ambiguous():
    item = {
        "category": "SMOKE_DUCT",
        "description": "Ong gio hut khoi 1800x500",
        "remark": "",
        "material": "GI",
        "width": 1800,
        "height": 500,
        "thickness": 0.75,
        "warnings": [],
    }
    memory = [
        {
            "category": "SMOKE_DUCT", "normalized_description": "ong gio hut khoi 1800x500",
            "material": "GI", "width": 1800, "height": 500, "thickness": 0.75,
            "quote_material_spec": "Ton ma kem 0.75mm boc MGO EI45", "quote_note": "",
        },
        {
            "category": "SMOKE_DUCT", "normalized_description": "ong gio hut khoi 1800x500",
            "material": "GI", "width": 1800, "height": 500, "thickness": 0.75,
            "quote_material_spec": "Ton ma kem 0.75mm boc MGO EI60", "quote_note": "",
        },
    ]

    result = resolve_material_specifications([item], memory_rows=memory, product_rules={})[0]

    assert result["quote_material_spec"] == ""
    assert result["material_spec_source"] == "needs_confirmation"
    assert result["warnings"]