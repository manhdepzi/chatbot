"""Build an output quotation by filling the Kaiyo template.

The builder copies ``report-template/kaiyo_quotation_template.xlsx`` and fills
the ``Mẫu - xây lắp`` sheet, preserving the template's styles, merges and
number formats. Column positions are detected from the template's own header
row so the same code works across Kaiyo template variants.
"""
from __future__ import annotations

import copy
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter

from .. import config

# header labels (lowercased) -> canonical field. Dimension tokens are matched
# exactly (or as a prefix) to avoid single-char false matches.
_DIM_LABELS: list[tuple[tuple[str, ...], str]] = [
    (("mã sp",), "ma_sp"),
    (("w1",), "w1"),
    (("h1",), "h1"),
    (("w2",), "w2"),
    (("h2",), "h2"),
    (("w3",), "w3"),
    (("h3",), "h3"),
    (("l/h",), "l"),
    (("r/d",), "r"),
    (("e",), "e"),
    (("diện tích",), "area"),
    (("giá tôn",), "gia_ton"),
    (("hệ số",), "he_so"),
]

_CORE_LABELS: list[tuple[tuple[str, ...], str]] = [
    (("stt",), "stt"),
    (("tên vật tư", "tên công tác", "sản phẩm", "tên"), "ten"),
    (("vật liệu chế tạo", "vật liệu", "quy cách"), "vat_lieu"),
    (("xuất xứ", "nhãn hiệu"), "xuat_xu"),
    (("đơn vị", "đvt"), "don_vi"),
    (("khối lượng", "số lượng"), "khoi_luong"),
    (("đơn giá",), "don_gia"),
    (("thành tiền",), "thanh_tien"),
    (("ghi chú",), "ghi_chu"),
]

_TOTAL_LABELS = {
    "tổng cộng giá trước thuế": "truoc_thue",
    "thuế gtgt 10%": "vat10",
    "thuế gtgt 8%": "vat8",
    "tổng cộng giá sau thuế": "sau_thue",
}


def _cell(v: Any) -> str:
    return "" if v is None else str(v).strip()


def _norm(s: str) -> str:
    return " ".join(s.split()).lower()


def _find_sheet(wb, name_hint: str = "Mẫu - xây lắp"):
    if name_hint in wb.sheetnames:
        return wb[name_hint]
    # fallback: pick the largest sheet (most rows)
    return max(wb.worksheets, key=lambda ws: ws.max_row)


def _find_header_row(ws) -> int | None:
    for r in range(1, min(ws.max_row, 40) + 1):
        cells = [_norm(_cell(ws.cell(r, c).value)) for c in range(1, ws.max_column + 1)]
        if "stt" in cells and any("tên" in x or "sản phẩm" in x for x in cells):
            return r
    return None


def _map_columns(ws, header_row: int) -> dict[str, int]:
    """Map canonical field -> 1-based column index from the header row.

    Dimension tokens (w1, h1, e, ...) match exactly to avoid single-char false
    positives; core fields match as substrings, with ``ten`` bound to the first
    name-like column and ``stt`` to the column literally labelled "stt".
    """
    mapping: dict[str, int] = {}
    taken: set[int] = set()

    # first pass: core columns, left-to-right
    for c in range(1, ws.max_column + 1):
        label = _norm(_cell(ws.cell(header_row, c).value))
        if not label:
            continue
        for keywords, canon in _CORE_LABELS:
            if canon in mapping:
                continue
            if canon == "stt":
                if label != "stt":
                    continue
            else:
                if not any(kw in label for kw in keywords):
                    continue
            mapping[canon] = c
            taken.add(c)
            break

    # second pass: dimension columns (exact match)
    for c in range(1, ws.max_column + 1):
        if c in taken:
            continue
        label = _norm(_cell(ws.cell(header_row, c).value))
        if not label:
            continue
        for keywords, canon in _DIM_LABELS:
            if canon in mapping:
                continue
            if canon == "area":
                # "DIỆN TÍCH /CÁI " normalises to "diện tích /cái"; match prefix
                if label.startswith("diện tích"):
                    mapping[canon] = c
                    taken.add(c)
                    break
            if label in keywords:
                mapping[canon] = c
                taken.add(c)
                break

    # the area table repeats the product name in a column immediately left of
    # Mã SP (its header label is merged K20:K21, so it is absent from the
    # header row itself). Bind it as ``ten_dim`` = the column just before Mã SP.
    if "ma_sp" in mapping and mapping["ma_sp"] > 1:
        mapping["ten_dim"] = mapping["ma_sp"] - 1

    return mapping


def _find_total_row(ws, start_row: int) -> int | None:
    for r in range(start_row, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            label = _norm(_cell(ws.cell(r, c).value))
            if "tổng cộng giá trước thuế" in label:
                return r
    return None


def _copy_style(src, dst):
    if src.has_style:
        dst.font = copy.copy(src.font)
        dst.fill = copy.copy(src.fill)
        dst.border = copy.copy(src.border)
        dst.alignment = copy.copy(src.alignment)
        dst.number_format = src.number_format
        dst.protection = copy.copy(src.protection)


def build_report(items: list[dict[str, Any]], header: dict[str, Any],
                 out_path: str | Path, totals: dict[str, float]) -> str:
    """Fill the template and save to ``out_path``. Returns the output path."""
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.load_workbook(config.TEMPLATE_FILE)
    ws = _find_sheet(wb)

    header_row = _find_header_row(ws)
    if header_row is None:
        raise RuntimeError("could not locate the item header row in template")

    col = _map_columns(ws, header_row)
    if "ten" not in col or "khoi_luong" not in col:
        raise RuntimeError("template header lacks name/quantity columns")

    first_data_row = header_row + 2  # template has a sub-header at +1
    total_row = _find_total_row(ws, first_data_row)

    # --- clear existing data rows between header and totals -----------------
    if total_row is None:
        total_row = first_data_row + len(items) + 6
    n_rows_needed = len(items)
    n_available = (total_row - first_data_row) if total_row else 0

    # reference style row (a clean data row): use header_row+2
    ref_row = header_row + 2

    # clear old data cells
    for r in range(first_data_row, total_row):
        for c in range(1, ws.max_column + 1):
            ws.cell(r, c).value = None

    # insert rows if needed (before totals)
    if n_rows_needed > n_available:
        extra = n_rows_needed - n_available
        ws.insert_rows(total_row, extra)
        total_row += extra

    # --- write items --------------------------------------------------------
    r = first_data_row
    for it in items:
        if it.get("is_section"):
            ws.cell(r, col["stt"]).value = it["ten"]
        else:
            ws.cell(r, col["stt"]).value = it.get("stt")
            ws.cell(r, col["ten"]).value = it.get("ten")
            if "vat_lieu" in col:
                ws.cell(r, col["vat_lieu"]).value = it.get("vat_lieu")
            if "xuat_xu" in col:
                ws.cell(r, col["xuat_xu"]).value = it.get("xuat_xu")
            if "don_vi" in col:
                ws.cell(r, col["don_vi"]).value = it.get("don_vi")
            if "khoi_luong" in col:
                ws.cell(r, col["khoi_luong"]).value = it.get("khoi_luong")
            if "don_gia" in col:
                ws.cell(r, col["don_gia"]).value = it.get("don_gia")
            if "thanh_tien" in col:
                ws.cell(r, col["thanh_tien"]).value = it.get("thanh_tien")
            if "ghi_chu" in col:
                ws.cell(r, col["ghi_chu"]).value = it.get("ghi_chu")
            if "ma_sp" in col:
                ws.cell(r, col["ma_sp"]).value = it.get("ma_sp")
            # the area table repeats the product name in the column just left
            # of Mã SP (golden fills it with "=+B", i.e. the left-hand name)
            if "ten_dim" in col:
                ws.cell(r, col["ten_dim"]).value = it.get("ten")
            # dimension columns
            for k, c in col.items():
                if k in ("w1", "h1", "w2", "h2", "w3", "h3", "l", "r", "e",
                         "area", "gia_ton", "he_so"):
                    ws.cell(r, c).value = it.get(k)
        # copy style from reference row so new/kept rows look like the template
        for c in range(1, ws.max_column + 1):
            _copy_style(ws.cell(ref_row, c), ws.cell(r, c))
        r += 1

    # --- totals -------------------------------------------------------------
    # Each total line is a separate row: its label sits in a text column and
    # the value goes into the "thành tiền" (amount) column.
    amt_col = col.get("thanh_tien", 8)
    total_row_map: dict[str, int] = {}
    for r in range(first_data_row, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            label = _norm(_cell(ws.cell(r, c).value))
            for lbl, canon in _TOTAL_LABELS.items():
                if lbl in label and canon not in total_row_map:
                    total_row_map[canon] = r
                    break
    for canon, r in total_row_map.items():
        if canon in totals:
            ws.cell(r, amt_col).value = totals[canon]

    # --- header fields ------------------------------------------------------
    _fill_header(ws, header)

    wb.save(out_path)
    return str(out_path)


def _fill_header(ws, header: dict[str, Any]):
    """Write quote header values into labelled cells."""
    if not header:
        return
    # map merged ranges: any merged cell -> its top-left anchor
    anchor = {}
    for mr in ws.merged_cells.ranges:
        min_r, min_c, max_r, max_c = mr.min_row, mr.min_col, mr.max_row, mr.max_col
        for rr in range(min_r, max_r + 1):
            for cc in range(min_c, max_c + 1):
                anchor[(rr, cc)] = (min_r, min_c)

    labels = {
        "kính gửi": "kinh_gui",
        "địa chỉ": "dia_chi",
        "người nhận": "nguoi_nhan",
        "dự án": "du_an",
        "số bg": "so_bg",
        "ngày bg": "ngay_bg",
    }
    for r in range(1, 30):
        for c in range(1, ws.max_column + 1):
            lbl = _norm(_cell(ws.cell(r, c).value))
            if not lbl:
                continue
            canon = None
            for k, v in labels.items():
                if lbl == k or lbl == k + ":":
                    canon = v
                    break
            if canon and header.get(canon):
                # write into the first empty cell to the right
                for cc in range(c + 1, min(c + 4, ws.max_column + 1)):
                    target = ws.cell(r, cc)
                    if (r, cc) in anchor:
                        tr, tc = anchor[(r, cc)]
                        target = ws.cell(tr, tc)
                    if target.value in (None, ""):
                        target.value = header[canon]
                        break
